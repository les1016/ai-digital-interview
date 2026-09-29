from __future__ import annotations

from datetime import datetime
from html import escape
from pathlib import Path
import re
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.pdfmetrics import registerFont
from reportlab.platypus import (
    CondPageBreak,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


NAVY = "244B48"
TEAL = "3B8179"
PALE_TEAL = "EAF5F2"
CORAL = "C96F50"
PALE_CORAL = "FFF1E9"
INK = "334B47"
MUTED = "6E817C"
LINE = "D9E7E3"
WHITE = "FFFFFF"


def _safe(value: Any) -> str:
    return escape(str(value or "")).replace("\n", "<br/>")


def _without_list_marker(value: Any) -> str:
    return re.sub(r"^\s*\d+(?:\.\d+)*[.、：:)]?\s*", "", str(value or "")).strip()


def _list(value: Any) -> list:
    return value if isinstance(value, list) else []


def _summary(report: dict) -> dict:
    return report.get("conversation_summary") or {}


def _pdf_styles():
    registerFont(UnicodeCIDFont("STSong-Light"))
    sample = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "ReportTitle", parent=sample["Title"], fontName="STSong-Light",
            fontSize=24, leading=31, textColor=colors.HexColor(f"#{NAVY}"),
            alignment=TA_CENTER, spaceAfter=8,
        ),
        "subtitle": ParagraphStyle(
            "ReportSubtitle", parent=sample["Normal"], fontName="STSong-Light",
            fontSize=10, leading=15, textColor=colors.HexColor(f"#{MUTED}"),
            alignment=TA_CENTER, spaceAfter=16,
        ),
        "h1": ParagraphStyle(
            "ReportH1", parent=sample["Heading1"], fontName="STSong-Light",
            fontSize=16, leading=22, textColor=colors.HexColor(f"#{NAVY}"),
            spaceBefore=14, spaceAfter=8,
        ),
        "h2": ParagraphStyle(
            "ReportH2", parent=sample["Heading2"], fontName="STSong-Light",
            fontSize=12, leading=18, textColor=colors.HexColor(f"#{TEAL}"),
            spaceBefore=9, spaceAfter=5,
        ),
        "body": ParagraphStyle(
            "ReportBody", parent=sample["BodyText"], fontName="STSong-Light",
            fontSize=10.5, leading=17, textColor=colors.HexColor(f"#{INK}"),
            spaceAfter=6,
        ),
        "small": ParagraphStyle(
            "ReportSmall", parent=sample["BodyText"], fontName="STSong-Light",
            fontSize=8.7, leading=14, textColor=colors.HexColor(f"#{MUTED}"),
            spaceAfter=4,
        ),
        "score": ParagraphStyle(
            "ReportScore", parent=sample["Title"], fontName="Helvetica-Bold",
            fontSize=34, leading=38, textColor=colors.HexColor(f"#{TEAL}"),
            alignment=TA_CENTER,
        ),
        "callout": ParagraphStyle(
            "ReportCallout", parent=sample["BodyText"], fontName="STSong-Light",
            fontSize=10.5, leading=17, textColor=colors.HexColor(f"#{INK}"),
            leftIndent=4, rightIndent=4,
        ),
        "label": ParagraphStyle(
            "ReportLabel", parent=sample["BodyText"], fontName="STSong-Light",
            fontSize=9, leading=14, textColor=colors.HexColor(f"#{CORAL}"),
            spaceAfter=3,
        ),
    }


def _pdf_footer(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor(f"#{LINE}"))
    canvas.line(doc.leftMargin, 0.58 * inch, LETTER[0] - doc.rightMargin, 0.58 * inch)
    canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
    canvas.setFont("STSong-Light", 8)
    canvas.drawString(doc.leftMargin, 0.38 * inch, "AI数字人面试 · 个人成长档案")
    canvas.drawRightString(LETTER[0] - doc.rightMargin, 0.38 * inch, f"第 {doc.page} 页")
    canvas.restoreState()


def build_pdf(
    output_path: Path,
    *,
    candidate_name: str,
    target_role: str,
    started_at: datetime,
    report: dict,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    styles = _pdf_styles()
    doc = SimpleDocTemplate(
        str(output_path), pagesize=LETTER,
        leftMargin=inch, rightMargin=inch, topMargin=0.78 * inch, bottomMargin=0.78 * inch,
        title=f"{candidate_name}的面试成长报告",
        author="AI数字人面试",
    )
    story = []
    story.append(Spacer(1, 0.15 * inch))
    story.append(Paragraph("AI数字人面试", styles["label"]))
    story.append(Paragraph("个人面试成长报告", styles["title"]))
    story.append(Paragraph(
        f"{_safe(candidate_name)}　|　{_safe(target_role or '综合岗位探索')}　|　{started_at.strftime('%Y年%m月%d日')}",
        styles["subtitle"],
    ))

    score_table = Table([
        [Paragraph(str(report.get("overall_score", 0)), styles["score"]),
         Paragraph(f"<b>整体表现</b><br/>{_safe(report.get('summary'))}", styles["body"])],
    ], colWidths=[1.25 * inch, 5.25 * inch])
    score_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{PALE_TEAL}")),
        ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor(f"#{LINE}")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 14),
        ("RIGHTPADDING", (0, 0), (-1, -1), 14),
        ("TOPPADDING", (0, 0), (-1, -1), 12),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
    ]))
    story.extend([score_table, Spacer(1, 10)])

    story.append(Paragraph("01　面试结果", styles["h1"]))
    story.append(Paragraph("你已经完成了一次有价值的练习", styles["h2"]))
    for item in _list(report.get("highlights")):
        story.append(Paragraph(f"✓　{_safe(item)}", styles["body"]))

    improvements = _list(report.get("improvements"))
    if improvements:
        improvement_rows = [[Paragraph("下一步提升建议", styles["label"])]]
        improvement_rows.extend([[Paragraph(f"•　{_safe(item)}", styles["callout"])]] for item in improvements)
        improvement_table = Table(improvement_rows, colWidths=[6.5 * inch])
        improvement_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{PALE_CORAL}")),
            ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor(f"#{CORAL}")),
            ("LEFTPADDING", (0, 0), (-1, -1), 14),
            ("RIGHTPADDING", (0, 0), (-1, -1), 14),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        story.extend([Spacer(1, 5), improvement_table])

    target_match = report.get("target_match")
    if isinstance(target_match, dict):
        story.append(Paragraph("目标岗位匹配度", styles["h2"]))
        story.append(Paragraph(
            f"<b>{_safe(target_match.get('target') or '目标岗位')}</b>　{target_match.get('score', 0)}%",
            styles["body"],
        ))
        story.append(Paragraph(_safe(target_match.get("evidence")), styles["small"]))
        for item in _list(target_match.get("matched_requirements")):
            story.append(Paragraph(f"✓　{_safe(item)}", styles["body"]))
        for item in _list(target_match.get("gaps")):
            story.append(Paragraph(f"待提升　{_safe(item)}", styles["callout"]))

    dimensions = report.get("dimensions") or {}
    if dimensions:
        story.append(Paragraph("能力维度", styles["h2"]))
        rows = [[
            Paragraph("维度", styles["small"]),
            Paragraph("得分", styles["small"]),
            Paragraph("证据与反馈", styles["small"]),
        ]]
        for name, item in dimensions.items():
            rows.append([
                Paragraph(_safe(name), styles["body"]),
                Paragraph(str(item.get("score", 0)), styles["body"]),
                Paragraph(_safe(item.get("evidence")), styles["small"]),
            ])
        table = Table(rows, colWidths=[1.25 * inch, 0.65 * inch, 4.6 * inch], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{LINE}")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (1, 1), (1, -1), "CENTER"),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        story.append(table)

    roles = _list(report.get("recommended_roles"))
    if roles:
        story.append(Paragraph("推荐探索方向", styles["h2"]))
        for item in roles:
            story.append(KeepTogether([
                Paragraph(f"<b>{_safe(item.get('role'))}</b>　匹配参考 {item.get('fit', 0)}%", styles["body"]),
                Paragraph(_safe(item.get("reason")), styles["small"]),
            ]))

    story.append(Paragraph("情绪与状态调节", styles["h2"]))
    for item in _list(report.get("emotion_guidance")):
        story.append(Paragraph(f"•　{_safe(item)}", styles["body"]))
    story.append(Paragraph("行动计划", styles["h2"]))
    for index, item in enumerate(_list(report.get("action_plan")), 1):
        story.append(Paragraph(f"{index}.　{_safe(_without_list_marker(item))}", styles["body"]))

    digest = _summary(report)
    story.extend([CondPageBreak(3.0 * inch), Paragraph("02　专属面试对话摘要", styles["h1"])])
    story.append(Paragraph(_safe(digest.get("overview") or "本摘要依据本次面试问答整理。"), styles["body"]))
    topics = _list(digest.get("key_topics"))
    if topics:
        story.append(Paragraph("本次讨论重点", styles["h2"]))
        story.append(Paragraph("　·　".join(_safe(item) for item in topics), styles["body"]))
    for index, item in enumerate(_list(digest.get("dialogue_digest")), 1):
        block = [
            Paragraph(f"第 {index} 组问答", styles["label"]),
            Paragraph(f"<b>问题：</b>{_safe(item.get('question'))}", styles["body"]),
            Paragraph(f"<b>回答摘要：</b>{_safe(item.get('answer_summary'))}", styles["body"]),
            Paragraph(f"<font color='#{CORAL}'><b>教练提示：</b>{_safe(item.get('coaching_note'))}</font>", styles["body"]),
        ]
        summary_table = Table([[block]], colWidths=[6.5 * inch])
        summary_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F7FAF9")),
            ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor(f"#{LINE}")),
            ("LEFTPADDING", (0, 0), (-1, -1), 12),
            ("RIGHTPADDING", (0, 0), (-1, -1), 12),
            ("TOPPADDING", (0, 0), (-1, -1), 9),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
        ]))
        story.extend([summary_table, Spacer(1, 8)])
    follow_up = _list(digest.get("follow_up_focus"))
    if follow_up:
        story.append(Paragraph("下一次练习重点", styles["h2"]))
        for item in follow_up:
            story.append(Paragraph(f"•　{_safe(item)}", styles["body"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph(_safe(report.get("disclaimer")), styles["small"]))
    doc.build(story, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)


def _set_run_font(run, *, size=None, color=None, bold=None, name="Microsoft YaHei"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def _shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            _cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _add_docx_bullet(doc, text, *, color=INK):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(str(text))
    _set_run_font(run, size=10.5, color=color)
    return paragraph


def _add_docx_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    paragraph.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    run = paragraph.add_run(text)
    _set_run_font(run, size=16 if level == 1 else 13, color=NAVY if level == 1 else TEAL, bold=True)
    return paragraph


def build_docx(
    output_path: Path,
    *,
    candidate_name: str,
    target_role: str,
    started_at: datetime,
    report: dict,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for style_name, size, color in (("Heading 1", 16, NAVY), ("Heading 2", 13, TEAL), ("Heading 3", 12, NAVY)):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("AI数字人面试　|　个人成长档案")
    _set_run_font(run, size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("AI数字人面试 · 结果与对话摘要合并文档")
    _set_run_font(run, size=8, color=MUTED)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(6)
    _set_run_font(kicker.add_run("AI数字人面试"), size=10, color=CORAL, bold=True)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    _set_run_font(title.add_run("个人面试成长报告"), size=25, color=NAVY, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    _set_run_font(
        subtitle.add_run(f"{candidate_name}　|　{target_role or '综合岗位探索'}　|　{started_at.strftime('%Y年%m月%d日')}"),
        size=10, color=MUTED,
    )

    score_table = doc.add_table(rows=1, cols=2)
    _table_widths(score_table, [1800, 7560])
    for cell in score_table.rows[0].cells:
        _shade(cell, PALE_TEAL)
    score_p = score_table.cell(0, 0).paragraphs[0]
    score_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(score_p.add_run(str(report.get("overall_score", 0))), size=30, color=TEAL, bold=True)
    summary_p = score_table.cell(0, 1).paragraphs[0]
    _set_run_font(summary_p.add_run("整体表现\n"), size=11, color=NAVY, bold=True)
    _set_run_font(summary_p.add_run(str(report.get("summary") or "")), size=10.5, color=INK)

    _add_docx_heading(doc, "01　面试结果", 1)
    lead = doc.add_paragraph()
    _set_run_font(lead.add_run("你已经完成了一次有价值的练习"), size=13, color=TEAL, bold=True)
    for item in _list(report.get("highlights")):
        _add_docx_bullet(doc, item)

    improvements = _list(report.get("improvements"))
    if improvements:
        callout = doc.add_table(rows=1 + len(improvements), cols=1)
        _table_widths(callout, [9360])
        for row in callout.rows:
            _shade(row.cells[0], PALE_CORAL)
        first = callout.cell(0, 0).paragraphs[0]
        _set_run_font(first.add_run("下一步提升建议"), size=11, color=CORAL, bold=True)
        for index, item in enumerate(improvements, 1):
            p = callout.cell(index, 0).paragraphs[0]
            _set_run_font(p.add_run(f"{index}. {item}"), size=10.5, color=INK)

    target_match = report.get("target_match")
    if isinstance(target_match, dict):
        _add_docx_heading(doc, "目标岗位匹配度", 2)
        p = doc.add_paragraph()
        _set_run_font(p.add_run(f"{target_match.get('score', 0)}%　"), size=20, color=TEAL, bold=True)
        _set_run_font(p.add_run(str(target_match.get("target") or "目标岗位")), size=11, color=NAVY, bold=True)
        evidence = doc.add_paragraph()
        _set_run_font(evidence.add_run(str(target_match.get("evidence") or "")), size=10.5, color=INK)
        for item in _list(target_match.get("matched_requirements")):
            _add_docx_bullet(doc, f"已匹配：{item}")
        for item in _list(target_match.get("gaps")):
            _add_docx_bullet(doc, f"待补齐：{item}")

    dimensions = report.get("dimensions") or {}
    if dimensions:
        _add_docx_heading(doc, "能力维度", 2)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        _table_widths(table, [1800, 900, 6660])
        for index, label in enumerate(("维度", "得分", "证据与反馈")):
            cell = table.rows[0].cells[index]
            _shade(cell, NAVY)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_run_font(p.add_run(label), size=9.5, color=WHITE, bold=True)
        for name, item in dimensions.items():
            cells = table.add_row().cells
            _table_widths(table, [1800, 900, 6660])
            for index, text in enumerate((name, item.get("score", 0), item.get("evidence", ""))):
                p = cells[index].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 1 else WD_ALIGN_PARAGRAPH.LEFT
                _set_run_font(p.add_run(str(text)), size=9.5, color=INK)

    roles = _list(report.get("recommended_roles"))
    if roles:
        _add_docx_heading(doc, "推荐探索方向", 2)
        for item in roles:
            p = doc.add_paragraph()
            _set_run_font(p.add_run(f"{item.get('role')}　匹配参考 {item.get('fit', 0)}%\n"), size=11, color=NAVY, bold=True)
            _set_run_font(p.add_run(str(item.get("reason") or "")), size=10.5, color=INK)

    _add_docx_heading(doc, "情绪与状态调节", 2)
    for item in _list(report.get("emotion_guidance")):
        _add_docx_bullet(doc, item)
    _add_docx_heading(doc, "行动计划", 2)
    for item in _list(report.get("action_plan")):
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(6)
        _set_run_font(paragraph.add_run(_without_list_marker(item)), size=10.5, color=INK)

    doc.add_section(WD_SECTION.NEW_PAGE)
    _add_docx_heading(doc, "02　专属面试对话摘要", 1)
    digest = _summary(report)
    overview = doc.add_paragraph()
    _set_run_font(overview.add_run(str(digest.get("overview") or "本摘要依据本次面试问答整理。")), size=11, color=INK)
    topics = _list(digest.get("key_topics"))
    if topics:
        _add_docx_heading(doc, "本次讨论重点", 2)
        for item in topics:
            _add_docx_bullet(doc, item)
    for index, item in enumerate(_list(digest.get("dialogue_digest")), 1):
        block = doc.add_table(rows=1, cols=1)
        _table_widths(block, [9360])
        _shade(block.cell(0, 0), "F7FAF9")
        p = block.cell(0, 0).paragraphs[0]
        _set_run_font(p.add_run(f"第 {index} 组问答\n"), size=10, color=CORAL, bold=True)
        _set_run_font(p.add_run(f"问题：{item.get('question', '')}\n"), size=10.5, color=NAVY, bold=True)
        _set_run_font(p.add_run(f"回答摘要：{item.get('answer_summary', '')}\n"), size=10.5, color=INK)
        _set_run_font(p.add_run(f"教练提示：{item.get('coaching_note', '')}"), size=10.5, color=CORAL, bold=True)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)
    follow_up = _list(digest.get("follow_up_focus"))
    if follow_up:
        _add_docx_heading(doc, "下一次练习重点", 2)
        for item in follow_up:
            _add_docx_bullet(doc, item, color=CORAL)
    disclaimer = doc.add_paragraph()
    disclaimer.paragraph_format.space_before = Pt(12)
    _set_run_font(disclaimer.add_run(str(report.get("disclaimer") or "")), size=8.5, color=MUTED)
    doc.save(str(output_path))


def build_export(
    output_path: Path,
    *,
    format_name: str,
    candidate_name: str,
    target_role: str,
    started_at: datetime,
    report: dict,
) -> Path:
    if format_name == "pdf":
        build_pdf(output_path, candidate_name=candidate_name, target_role=target_role, started_at=started_at, report=report)
        return output_path
    if format_name == "docx":
        build_docx(output_path, candidate_name=candidate_name, target_role=target_role, started_at=started_at, report=report)
        return output_path
    raise ValueError("unsupported document format")
